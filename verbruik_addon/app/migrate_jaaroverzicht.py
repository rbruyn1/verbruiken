"""Eenmalige migratie van de jaaroverzicht-cijfers uit de Excel.

Deze stonden niet in een nette tabel, maar als losse cellen tussen de
maandrijen: telkens in de rij van de kalendermaand waarin de eigenaar ze
bijhield (jan=elektriciteit, feb=opladen zon, apr=uitgespaard, mei=batterij,
okt=gas bij Tienen, jul/aug=verschil -1 jaar). De celcoördinaten hieronder
zijn expliciet vastgelegd na verificatie tegen de al-gemigreerde
maanddata (sommen kloppen tot op afrondingsniveau).

Twee vondsten bij Tienen zijn BEWUST weggelaten:
- Cellen AC16:AF17 (rij 2021-07) en AC28:AF29 (rij 2022-07) dragen allebei
  het label "verschil ... 2022 - 2023", wat niet bij hun eigen jaar past
  (zou "2020-2021" resp. "2021-2022" moeten zijn) — duidelijk een
  stale/foutief gekopieerde cel uit de Excel zelf, geen betrouwbare data.
  Voor 2021 en 2022 is er dus geen "verschil -1 jaar"-data.

Bij Binkom draagt de verschil-cel het label "verschil elektriciteit
2024 - 2026" (spanning 2 jaar, i.p.v. het gebruikelijke 1-jaars-patroon)
— letterlijk overgenomen zoals in de Excel staat, niet gecorrigeerd.

Gebruik: python migrate_jaaroverzicht.py /pad/naar/nettarieven_2023.xlsx
"""
import sys
from pathlib import Path
import openpyxl
from models import get_connection, init_db


def cell(ws, coord):
    v = ws[coord].value
    if isinstance(v, str):
        return None  # tekstlabel, geen waarde
    return v


def build_tienen(ws):
    return {
        2021: {
            "jaarverbruik_elektriciteit_kwh": cell(ws, "AC10"),
            "jaarverbruik_elektriciteit_kost_eur": cell(ws, "AD10"),
        },
        2022: {
            "jaarverbruik_elektriciteit_kwh": cell(ws, "AC22"),
            "jaarverbruik_elektriciteit_kost_eur": cell(ws, "AD22"),
            "jaarverbruik_gas": cell(ws, "AC31"),
            "jaarverbruik_gas_eenheid": "kWh",
            "jaarverbruik_gas_kost_eur": cell(ws, "AD31"),
        },
        2023: {
            "jaarverbruik_elektriciteit_kwh": cell(ws, "AC34"),
            "jaarverbruik_elektriciteit_kost_eur": cell(ws, "AD34"),
            "jaarverbruik_gas": cell(ws, "AC43"),
            "jaarverbruik_gas_eenheid": "kWh",
            "jaarverbruik_gas_kost_eur": cell(ws, "AD43"),
            "uitgespaard_zonnepanelen_eur": cell(ws, "AE37"),
            "verschil_gas_vorig_jaar_m3": cell(ws, "AD40"),
            "verschil_gas_vorig_jaar_kwh": cell(ws, "AF40"),
            "verschil_elektriciteit_vorig_jaar_kwh": cell(ws, "AD41"),
            "verschil_elektriciteit_vorig_jaar_eur": cell(ws, "AF41"),
        },
        2024: {
            "jaarverbruik_elektriciteit_kwh": cell(ws, "AC46"),
            "jaarverbruik_elektriciteit_kost_eur": cell(ws, "AD46"),
            "opladen_zon_kwh": cell(ws, "AC47"),
            "opladen_zon_kost_eur": cell(ws, "AD47"),
            "jaarverbruik_gas": cell(ws, "AC55"),
            "jaarverbruik_gas_eenheid": "kWh",
            "jaarverbruik_gas_kost_eur": cell(ws, "AD55"),
            "uitgespaard_zonnepanelen_eur": cell(ws, "AE49"),
            "verschil_gas_vorig_jaar_m3": cell(ws, "AD52"),
            "verschil_gas_vorig_jaar_kwh": cell(ws, "AF52"),
            "verschil_elektriciteit_vorig_jaar_kwh": cell(ws, "AD53"),
            "verschil_elektriciteit_vorig_jaar_eur": cell(ws, "AF53"),
        },
        2025: {
            "jaarverbruik_elektriciteit_kwh": cell(ws, "AC58"),
            "jaarverbruik_elektriciteit_kost_eur": cell(ws, "AD58"),
            "opladen_zon_kwh": cell(ws, "AC59"),
            "opladen_zon_kost_eur": cell(ws, "AD59"),
            "jaarverbruik_gas": cell(ws, "AC67"),
            "jaarverbruik_gas_eenheid": "kWh",
            "jaarverbruik_gas_kost_eur": cell(ws, "AD67"),
            "uitgespaard_zonnepanelen_eur": cell(ws, "AE61"),
            "verschil_gas_vorig_jaar_m3": cell(ws, "AD64"),
            "verschil_gas_vorig_jaar_kwh": cell(ws, "AF64"),
            "verschil_elektriciteit_vorig_jaar_kwh": cell(ws, "AD65"),
            "verschil_elektriciteit_vorig_jaar_eur": cell(ws, "AF65"),
        },
        2026: {
            "jaarverbruik_elektriciteit_kwh": cell(ws, "AC70"),
            "jaarverbruik_elektriciteit_kost_eur": cell(ws, "AD70"),
            "opladen_zon_kwh": cell(ws, "AC71"),
            "opladen_zon_kost_eur": cell(ws, "AD71"),
            "jaarverbruik_gas": cell(ws, "AC79"),
            "jaarverbruik_gas_eenheid": "kWh",
            "jaarverbruik_gas_kost_eur": cell(ws, "AD79"),
            "uitgespaard_zonnepanelen_eur": cell(ws, "AE73"),
            "batterij_gebruik_kwh": cell(ws, "AD74"),
            "batterij_gebruik_kost_eur": cell(ws, "AE74"),
            "verschil_gas_vorig_jaar_m3": cell(ws, "AD76"),
            "verschil_gas_vorig_jaar_kwh": cell(ws, "AF76"),
            "verschil_elektriciteit_vorig_jaar_kwh": cell(ws, "AD77"),
            "verschil_elektriciteit_vorig_jaar_eur": cell(ws, "AF77"),
        },
    }


def build_binkom(ws):
    return {
        2025: {
            "jaarverbruik_elektriciteit_kwh": cell(ws, "Y3"),
            "jaarverbruik_elektriciteit_kost_eur": cell(ws, "Z3"),
            "opladen_zon_kwh": cell(ws, "Y4"),
            "opladen_zon_kost_eur": cell(ws, "Z4"),
            "uitgespaard_zonnepanelen_eur": cell(ws, "AA6"),
        },
        2026: {
            "jaarverbruik_elektriciteit_kwh": cell(ws, "Y15"),
            "jaarverbruik_elektriciteit_kost_eur": cell(ws, "Z15"),
            "opladen_zon_kwh": cell(ws, "Y16"),
            "opladen_zon_kost_eur": cell(ws, "Z16"),
            "uitgespaard_zonnepanelen_eur": cell(ws, "AA18"),
            "batterij_gebruik_kwh": cell(ws, "Z19"),
            "batterij_gebruik_kost_eur": cell(ws, "AA19"),
            "verschil_elektriciteit_vorig_jaar_kwh": cell(ws, "Z22"),
            "verschil_elektriciteit_vorig_jaar_eur": cell(ws, "AB22"),
        },
    }


def migrate(xlsx_path):
    init_db()
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    conn = get_connection()

    data = {
        "Tienen": build_tienen(wb["verbruiken Tienen"]),
        "Binkom": build_binkom(wb["verbruiken Binkom"]),
    }

    ingevoerd = 0
    for woning, jaren in data.items():
        for jaar, velden in jaren.items():
            record = {"woning": woning, "jaar": jaar, **velden}
            kolommen = ", ".join(record.keys())
            placeholders = ", ".join(f":{k}" for k in record)
            update = ", ".join(
                f"{k}=excluded.{k}" for k in record if k not in ("woning", "jaar")
            )
            conn.execute(
                f"""INSERT INTO jaaroverzicht ({kolommen}) VALUES ({placeholders})
                    ON CONFLICT(woning, jaar) DO UPDATE SET {update}""",
                record,
            )
            ingevoerd += 1

    conn.commit()
    conn.close()
    print(f"Klaar: {ingevoerd} jaaroverzicht-records ingeladen.")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Gebruik: python migrate_jaaroverzicht.py <pad naar xlsx>")
        sys.exit(1)
    migrate(Path(sys.argv[1]))
